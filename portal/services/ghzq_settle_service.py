"""国海证券融资融券对账单 xlsx：提取「1.1当前资产情况」「2、负债情况」并合并为单一 metrics（英文字段）。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# 表头中文 → 英文字段名（1.1 资产 + 2 负债 共用；两表各取本表存在的列）
_GHZQ_HEADER_CN_TO_EN: dict[str, str] = {
    "币种": "currency",
    "资金账号": "fund_account_id",
    "总资产": "total_assets",
    "净资产": "net_assets",
    "资金余额": "cash_balance",
    "冻结资金": "frozen_cash",
    "证券市值": "securities_market_value",
    "期初余额": "opening_balance",
    "保证金可用": "margin_available",
    "可取金额": "withdrawable_amount",
    "融资余额": "financing_balance",
    "未了结融资利息": "outstanding_financing_interest",
    "融资费用": "financing_fee",
    "融资保证金": "financing_margin",
    "融券市值": "short_market_value",
    "融券费用": "short_fee",
    "未了结融券利息": "outstanding_short_interest",
    "其他负债": "other_liabilities",
    "未了结其他负债利息": "outstanding_other_liabilities_interest",
    "融券保证金": "short_margin",
    "待扣收": "pending_deduction",
    "转融通成本费用": "ref_cost_fee",
    "负债合计": "total_liabilities",
}


def _norm_header(s: object) -> str:
    t = str(s or "").strip().replace("\n", "")
    return re.sub(r"\s+", "", t)


def _map_row_to_en(headers: list[Any], values: list[Any], mapping: dict[str, str]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    n = min(len(headers), len(values))
    for i in range(n):
        h = headers[i]
        key_cn = _norm_header(h)
        if not key_cn:
            continue
        en = mapping.get(key_cn)
        if en is None:
            for cn, ek in mapping.items():
                if cn and cn in key_cn:
                    en = ek
                    break
        if en is None:
            continue
        out[en] = _coerce_cell(values[i])
    return out


def _coerce_cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int) and not isinstance(v, bool):
        return int(v)
    if isinstance(v, float):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s) if "." in s or "e" in s.lower() else int(s)
    except ValueError:
        return str(v)


def _find_table_rows_after_keyword(file_path: Path, keyword: str) -> list[list[Any]] | None:
    wb = load_workbook(file_path, data_only=True)
    try:
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            found_r: int | None = None
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and keyword in str(cell.value):
                        found_r = cell.row
                        break
                if found_r is not None:
                    break
            if found_r is None:
                continue
            start = found_r + 1
            out: list[list[Any]] = []
            for row in ws.iter_rows(min_row=start, max_row=start + 24, values_only=True):
                vals = list(row)
                while vals and vals[-1] is None:
                    vals.pop()
                if not any(v is not None and str(v).strip() for v in vals):
                    if out:
                        break
                    continue
                out.append(vals)
                if len(out) >= 22:
                    break
            if len(out) >= 2:
                return out
        return None
    finally:
        wb.close()


def extract_ghzq_statement_from_xlsx(path: Path) -> dict[str, Any]:
    """解析对账单 xlsx，将 1.1 资产与 2 负债表合并为单一 metrics（英文字段）。"""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(str(path))

    asset_rows = _find_table_rows_after_keyword(path, "1.1当前资产情况")
    liability_rows = _find_table_rows_after_keyword(path, "2、负债情况")
    if not asset_rows or len(asset_rows) < 2:
        raise RuntimeError("未找到「1.1当前资产情况」表格或数据行不完整。")
    if not liability_rows or len(liability_rows) < 2:
        raise RuntimeError("未找到「2、负债情况」表格或数据行不完整。")

    asset = _map_row_to_en(asset_rows[0], asset_rows[1], _GHZQ_HEADER_CN_TO_EN)
    liability = _map_row_to_en(liability_rows[0], liability_rows[1], _GHZQ_HEADER_CN_TO_EN)

    metrics: dict[str, Any] = {}
    metrics.update(asset)
    metrics.update(liability)

    fund_account_id = ""
    if metrics.get("fund_account_id") is not None:
        fund_account_id = str(metrics["fund_account_id"]).strip()

    return {
        "fund_account_id": fund_account_id,
        "metrics": metrics,
    }
